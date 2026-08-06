from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User, Group
from django.contrib.auth.hashers import make_password
from django.contrib import messages
from django.http import HttpResponse

from .models import WorkRecord
from .forms import RegisterForm

from datetime import date
import json
import re

from django.db.models import Count
from .models import WorkRecord, UserProfile

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    Image
)

from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
import os
from django.conf import settings
from reportlab.lib.pagesizes import A4, landscape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
# ==========================================
# LOGIN
# ==========================================
def login_user(request):

    if request.method == "POST":

        username = request.POST.get("username")
        password = request.POST.get("password")

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:
            login(request, user)
            return redirect("/")

        messages.error(request, "Invalid Username or Password")

    return render(request, "calculator/login.html")


# ==========================================
# LOGOUT
# ==========================================
def logout_user(request):
    logout(request)
    return redirect("/login/")
# ==========================================
# HOME
# ==========================================
@login_required(login_url="/login/")
def index(request):

    if request.method == "POST":

        work_date = request.POST.get("date")
        operator = request.POST.get("operator")
        machine = request.POST.get("machine")
        site = request.POST.get("site")
        start_time = request.POST.get("start_time")
        end_time = request.POST.get("end_time")
        total_time = request.POST.get("total_time")
        break_time = request.POST.get("break_time")
        working_time = request.POST.get("working_time")

        WorkRecord.objects.create(
            user=request.user,
            date=work_date,
            operator=operator,
            machine=machine,
            site=site,
            start_time=start_time,
            end_time=end_time,
            total_time=total_time,
            break_time=break_time,
            working_time=working_time,
        )

        messages.success(request, "✅ Record Saved Successfully!")
        return redirect("/")

    return render(request, "calculator/index.html")


import re

@login_required(login_url="/login/")
def history(request):

    records = WorkRecord.objects.filter(
        user=request.user
    ).order_by("-created_at")

    search = request.GET.get("search", "").strip()
    from_date = request.GET.get("from_date", "")
    to_date = request.GET.get("to_date", "")

    if search:
        records = records.filter(operator__icontains=search)

    if from_date:
        records = records.filter(date__gte=from_date)

    if to_date:
        records = records.filter(date__lte=to_date)

    # Total Records
    total_records = records.count()

    total_work_minutes = 0
    total_break_minutes = 0

    chart_labels = []
    chart_data = []

    for r in records:

        m = re.findall(r"\d+", r.working_time)
        if len(m) >= 2:
            h = int(m[0])
            mn = int(m[1])

            total_work_minutes += h * 60 + mn

            chart_labels.append(r.operator)
            chart_data.append(round((h * 60 + mn) / 60, 2))

        b = re.findall(r"\d+", r.break_time)
        if len(b) >= 2:
            bh = int(b[0])
            bm = int(b[1])

            total_break_minutes += bh * 60 + bm

    context = {
        "records": records,

        "total_records": total_records,

        "search": search,
        "from_date": from_date,
        "to_date": to_date,

        "total_work_hours": total_work_minutes // 60,
        "total_work_minutes": total_work_minutes % 60,

        "total_break_hours": total_break_minutes // 60,
        "total_break_minutes": total_break_minutes % 60,

        "chart_labels": chart_labels,
        "chart_data": chart_data,
    }

    return render(request, "calculator/history.html", context)
# ==========================================
# DELETE RECORD
# ==========================================
@login_required(login_url="/login/")
def delete_record(request, id):

    record = get_object_or_404(
        WorkRecord,
        id=id,
        user=request.user
    )

    record.delete()

    messages.success(request, "🗑 Record Deleted Successfully!")

    return redirect("/history/")


# ==========================================
# EDIT RECORD
# ==========================================
@login_required(login_url="/login/")
def edit_record(request, id):

    record = get_object_or_404(
        WorkRecord,
        id=id,
        user=request.user
    )

    if request.method == "POST":

        record.date = request.POST.get("date")
        record.operator = request.POST.get("operator")
        record.machine = request.POST.get("machine")
        record.site = request.POST.get("site")
        record.start_time = request.POST.get("start_time")
        record.end_time = request.POST.get("end_time")
        record.total_time = request.POST.get("total_time")
        record.break_time = request.POST.get("break_time")
        record.working_time = request.POST.get("working_time")

        record.save()

        messages.success(request, "✅ Record Updated Successfully!")

        return redirect("/history/")

    return render(request, "calculator/edit.html", {
        "record": record
    })
# ==========================================
# EXPORT PDF
# ==========================================
@login_required(login_url="/login/")
def export_pdf(request):

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="JCB_Work_Report.pdf"'

    doc = SimpleDocTemplate(
    response,
    pagesize=landscape(A4),
    rightMargin=20,
    leftMargin=20,
    topMargin=20,
    bottomMargin=20,
)

    styles = getSampleStyleSheet()

    title = styles["Title"]
    title.alignment = 1

    heading = styles["Heading2"]
    heading.alignment = 1

    normal = styles["Normal"]

    elements = []

    logo_path = os.path.join(
        settings.BASE_DIR,
        "calculator",
        "static",
        "images",
        "jcb_logo.png"
    )

    if os.path.exists(logo_path):
        logo = Image(logo_path, width=1.2*inch, height=1.2*inch)
        logo.hAlign = "CENTER"
        elements.append(logo)

    elements.append(Spacer(1, 10))

    elements.append(
        Paragraph(
            "<b>JCB WORK TIME REPORT</b>",
            title
        )
    )

    elements.append(
        Paragraph(
            "Professional Work Record Management System",
            heading
        )
    )

    elements.append(Spacer(1, 15))

    elements.append(
        Paragraph(
            f"<b>Generated By :</b> {request.user.username}",
            normal
        )
    )

    from datetime import datetime

    elements.append(
        Paragraph(
            f"<b>Date :</b> {datetime.now().strftime('%d-%m-%Y %H:%M')}",
            normal
        )
    )

    elements.append(Spacer(1, 20))

    data = [[
    "ID",
    "Date",
    "Operator",
    "Machine",
    "Site",
    "Start Time",
    "End Time",
    "Break Time",
    "Working Time",
    "Total Time",
    ]]

    records = WorkRecord.objects.filter(
        user=request.user
    ).order_by("-id")

    for r in records:
        data.append([
            str(r.id),
            str(r.date),
            r.operator,
            r.machine,
            r.site,
            str(r.start_time),
            str(r.end_time),
            r.break_time,
            r.working_time,
            r.total_time,
        ])
    table = Table(
    data,
    colWidths=[
        30,   # ID
        70,   # Date
        80,   # Operator
        80,   # Machine
        80,   # Site
        55,   # Start Time
        55,   # End Time
        70,   # Break Time
        70,   # Working Time
        70,   # Total Time
    ]
)

    table.setStyle(TableStyle([

    ("BACKGROUND", (0,0), (-1,0), HexColor("#198754")),
    ("TEXTCOLOR", (0,0), (-1,0), colors.white),

    ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
    ("FONTNAME", (0,1), (-1,-1), "Helvetica"),

    ("FONTSIZE", (0,0), (-1,-1),7 ),

    ("ALIGN", (0,0), (-1,-1), "CENTER"),
    ("VALIGN", (0,0), (-1,-1), "MIDDLE"),

    ("GRID", (0,0), (-1,-1), 0.5, colors.grey),

    ("BACKGROUND", (0,1), (-1,-1), colors.whitesmoke),

    ("BOTTOMPADDING", (0,0), (-1,-1), 6),
    ("TOPPADDING", (0,0), (-1,-1), 6),
    ("ROWBACKGROUNDS", (0,1), (-1,-1),
 [colors.white, colors.HexColor("#F4F6F8")]),
]))
    elements.append(table)
    elements.append(Spacer(1, 20))

    elements.append(
    Paragraph(
    "<para align='center'><font color='grey'>Generated from JCB Work Time Management System</font></para>",
    normal
    )
    )
    doc.build(elements)

    return response


# ==========================================
# EXPORT EXCEL
# ==========================================
@login_required(login_url="/login/")
def export_excel(request):

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "JCB Work Records"

    # ==========================
    # TITLE
    # ==========================

    sheet.merge_cells("A1:J1")

    title = sheet["A1"]
    title.value = "JCB WORK TIME REPORT"
    title.font = Font(bold=True, size=18, color="FFFFFF")
    title.fill = PatternFill(fill_type="solid", fgColor="198754")
    title.alignment = Alignment(horizontal="center")

    # ==========================
    # GENERATED BY
    # ==========================

    sheet.merge_cells("A2:J2")

    info = sheet["A2"]
    info.value = f"Generated By : {request.user.username}"
    info.font = Font(size=10, italic=True)
    info.alignment = Alignment(horizontal="right")

    # ==========================
    # TABLE HEADER
    # ==========================

    headers = [
        "ID",
        "Date",
        "Operator",
        "Machine",
        "Site",
        "Start Time",
        "End Time",
        "Total Time",
        "Break Time",
        "Working Time",
    ]

    for col, value in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=col)
        cell.value = value
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="198754"
        )
        cell.alignment = Alignment(horizontal="center")

    # ==========================
    # DATA
    # ==========================

    row = 4

    for record in WorkRecord.objects.filter(user=request.user):

        sheet.cell(row=row, column=1).value = record.id
        sheet.cell(row=row, column=2).value = str(record.date)
        sheet.cell(row=row, column=3).value = record.operator
        sheet.cell(row=row, column=4).value = record.machine
        sheet.cell(row=row, column=5).value = record.site
        sheet.cell(row=row, column=6).value = str(record.start_time)
        sheet.cell(row=row, column=7).value = str(record.end_time)
        sheet.cell(row=row, column=8).value = record.total_time
        sheet.cell(row=row, column=9).value = record.break_time
        sheet.cell(row=row, column=10).value = record.working_time

        row += 1

    # ==========================
    # BORDER
    # ==========================

    thin = Side(style="thin")

    for rows in sheet.iter_rows():
        for cell in rows:
            cell.border = Border(
                left=thin,
                right=thin,
                top=thin,
                bottom=thin
            )

    # ==========================
    # AUTO WIDTH
    # ==========================

    for column in sheet.columns:

        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass

        sheet.column_dimensions[column_letter].width = max_length + 4

    # ==========================
    # DOWNLOAD
    # ==========================

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="JCB_Work_Records.xlsx"'
    )

    workbook.save(response)

    return response

# ==========================================
# DASHBOARD
# ==========================================
@login_required(login_url="/login/")
def dashboard(request):

    records = WorkRecord.objects.filter(user=request.user)

    total_records = records.count()

    total_operators = (
        records.values("operator")
        .distinct()
        .count()
    )

    total_machines = (
        records.values("machine")
        .distinct()
        .count()
    )

    today_records = records.filter(
        date=date.today()
    ).count()

    latest_records = records.order_by("-id")[:5]

    chart = (
        records.values("operator")
        .annotate(total=Count("id"))
        .order_by("operator")
    )

    labels = []
    data = []

    for item in chart:
        labels.append(item["operator"])
        data.append(item["total"])

    context = {
        "total_records": total_records,
        "total_operators": total_operators,
        "total_machines": total_machines,
        "today_records": today_records,
        "latest_records": latest_records,
        "labels": json.dumps(labels),
        "data": json.dumps(data),
    }

    return render(
        request,
        "calculator/dashboard.html",
        context
    )
# ==========================================
# USERS
# ==========================================
@login_required(login_url="/login/")
def users(request):

    users_list = User.objects.all().order_by("id")

    return render(
        request,
        "calculator/users.html",
        {
            "users": users_list
        }
    )


# ==========================================
# ADD USER
# ==========================================
@login_required(login_url="/login/")
def add_user(request):

    if request.method == "POST":

        username = request.POST.get("username")
        email = request.POST.get("email")
        password = request.POST.get("password")

        if User.objects.filter(username=username).exists():
            messages.error(request, "Username already exists.")
            return redirect("/users/")

        User.objects.create_user(
            username=username,
            email=email,
            password=password
        )

        messages.success(request, "User created successfully.")
        return redirect("/users/")

    return render(request, "calculator/add_user.html")


# ==========================================
# EDIT USER
# ==========================================
@login_required(login_url="/login/")
def edit_user(request, id):

    user = get_object_or_404(User, id=id)

    if request.method == "POST":
        user.username = request.POST.get("username")
        user.email = request.POST.get("email")
        user.save()

        messages.success(request, "User updated successfully.")
        return redirect("/users/")

    return render(request, "calculator/edit_user.html", {"user": user})


# ==========================================
# DELETE USER
# ==========================================
@login_required(login_url="/login/")
def delete_user(request, id):

    user = get_object_or_404(User, id=id)
    user.delete()

    messages.success(request, "User deleted successfully.")
    return redirect("/users/")


# ==========================================
# RESET PASSWORD
# ==========================================
@login_required(login_url="/login/")
def reset_password(request, id):

    user = get_object_or_404(User, id=id)

    if request.method == "POST":
        password = request.POST.get("password")
        user.set_password(password)
        user.save()

        messages.success(request, "Password reset successfully.")
        return redirect("/users/")

    return render(request, "calculator/reset_password.html", {"user": user})


# ==========================================
# PROFILE
# ==========================================
@login_required(login_url="/login/")
def profile(request):
    role = "Admin" if request.user.is_superuser else "User"

    profile, created = UserProfile.objects.get_or_create(user=request.user)

    return render(
    request,
    "calculator/profile.html",
    {
        "user": request.user,
        "role": role,
        "profile": profile,
    }
)


# ==========================================
# EDIT PROFILE
# ==========================================
@login_required(login_url="/login/")
def edit_profile(request):

    profile, created = UserProfile.objects.get_or_create(
        user=request.user
    )

    if request.method == "POST":

        request.user.username = request.POST.get("username")
        request.user.email = request.POST.get("email")

        request.user.save()

        if request.FILES.get("photo"):
            profile.profile_photo = request.FILES.get("photo")
            profile.save()

        messages.success(request, "Profile updated successfully.")
        return redirect("/profile/")

    return render(
        request,
        "calculator/edit_profile.html",
        {
            "profile": profile
        }
    )
# ==========================================
# CHANGE PASSWORD
# ==========================================
@login_required(login_url="/login/")
def change_password(request):

    if request.method == "POST":

        current = request.POST.get("current")
        new = request.POST.get("new")
        confirm = request.POST.get("confirm")

        if not request.user.check_password(current):
            messages.error(request, "Current password is incorrect.")
            return redirect("/change-password/")

        if new != confirm:
            messages.error(request, "Passwords do not match.")
            return redirect("/change-password/")

        request.user.set_password(new)
        request.user.save()
        update_session_auth_hash(request, request.user)

        messages.success(request, "Password changed successfully.")
        return redirect("/profile/")

    return render(request, "calculator/change_password.html")


# ==========================================
# ABOUT
# ==========================================
def about(request):
    return render(request, "calculator/about.html")


# ==========================================
# REGISTER
# ==========================================
def register(request):

    if request.method == "POST":
        form = RegisterForm(request.POST)

        if form.is_valid():
            form.save()
            messages.success(request, "Registration successful. Please login.")
            return redirect("login")

    else:
        form = RegisterForm()

    return render(
        request,
        "calculator/register.html",
        {
            "form": form
        }
    )